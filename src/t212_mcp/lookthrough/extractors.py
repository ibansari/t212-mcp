"""Execute recipes: fetch the source, parse it, and return normalised fund holdings."""

import csv
import io
import json
import re
from contextlib import asynccontextmanager
from datetime import date, datetime, timezone
from typing import Any
from urllib.parse import urljoin

import httpx
from pydantic import BaseModel

from .countries import normalize_country
from .recipes import Recipe, fill

USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/130.0.0.0 Safari/537.36"
)


class ExtractionError(Exception):
    pass


class Holding(BaseModel):
    name: str
    weight_pct: float
    isin: str | None = None
    country: str | None = None
    sector: str | None = None


class FundHoldings(BaseModel):
    isin: str
    as_of: str | None
    coverage: str
    holdings: list[Holding]
    countries: dict[str, float] = {}
    sectors: dict[str, float] = {}
    source_url: str
    # First ~1.5k chars of the source's preamble/header area; used to confirm the data is for this fund.
    source_excerpt: str = ""
    recipe_id: str | None = None
    fetched_at: str = ""


# ---------------------------------------------------------------- fetching


def normalize_download_url(url: str) -> str:
    """Turn share links for Google Sheets / Drive into direct downloads."""
    m = re.search(r"docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]{20,})", url)
    if m:
        return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"
    m = re.search(r"drive\.google\.com/(?:file/d/|open\?id=)([A-Za-z0-9_-]{20,})", url)
    if m:
        return f"https://drive.google.com/uc?export=download&id={m.group(1)}"
    return url


def http_client() -> httpx.AsyncClient:
    return httpx.AsyncClient(
        headers={"User-Agent": USER_AGENT, "Accept-Language": "en-GB,en;q=0.9"},
        follow_redirects=True,
        timeout=45.0,
    )


async def http_get(url: str) -> httpx.Response:
    async with http_client() as c:
        resp = await c.get(url)
    if resp.is_error:
        raise ExtractionError(f"GET {url} -> HTTP {resp.status_code}")
    return resp


class BrowserSession:
    """A plain headless Chromium (no stealth). Pages are fetched like a normal visitor would."""

    def __init__(self, page):
        self.page = page

    async def goto(self, url: str) -> int | None:
        resp = await self.page.goto(url, wait_until="domcontentloaded", timeout=45000)
        await self.page.wait_for_timeout(1500)
        return resp.status if resp else None

    async def fetch_in_page(self, url: str) -> tuple[int, str]:
        return await self.page.evaluate(
            "async u => { const r = await fetch(u, {headers: {Accept: 'application/json, */*'}});"
            " return [r.status, await r.text()]; }",
            url,
        )


@asynccontextmanager
async def browser_session():
    from playwright.async_api import async_playwright

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        try:
            page = await browser.new_page(user_agent=USER_AGENT)
            yield BrowserSession(page)
        finally:
            await browser.close()


# ---------------------------------------------------------------- parsing helpers


def parse_table(content: bytes, fmt: str) -> list[list[Any]]:
    if fmt == "xls":
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        return rows
    if fmt == "xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    if fmt == "csv":
        text = content.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    raise ExtractionError(f"Unsupported table format {fmt}")


def sniff_format(content: bytes, content_type: str = "", url: str = "") -> str:
    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    if content[:2] == b"PK":
        return "xlsx"
    head = content[:200].lstrip()
    if head[:1] in (b"{", b"["):
        return "json"
    if "csv" in content_type or url.lower().split("?")[0].endswith(".csv"):
        return "csv"
    if "html" in content_type or head[:1] == b"<":
        return "html"
    return "csv"


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def find_header_row(rows: list[list[Any]], marker: str | None, columns_needed: list[str]) -> int:
    wanted = {c.lower() for c in columns_needed}
    for i, row in enumerate(rows[:200]):
        cells = {_cell_str(v).lower() for v in row}
        if marker and marker.lower() in cells:
            return i
        if not marker and wanted <= cells:
            return i
    raise ExtractionError(f"Header row not found (looked for {marker or sorted(wanted)})")


def to_float(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("%", "")
    try:
        return float(s)
    except ValueError:
        return None


DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%Y%m%d", "%d.%m.%Y")


def to_date(v: Any) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    m = re.match(r"\d{4}-\d{2}-\d{2}", s)
    if m:
        return m.group(0)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def dot_get(obj: Any, path: str | None) -> Any:
    if not path:
        return obj
    for part in path.split("."):
        if isinstance(obj, list):
            obj = obj[int(part)] if part.isdigit() and int(part) < len(obj) else None
        elif isinstance(obj, dict):
            obj = obj.get(part)
        else:
            return None
        if obj is None:
            return None
    return obj


def _ci_get(d: dict, key: str | None) -> Any:
    if key is None:
        return None
    if key in d:
        return d[key]
    lk = key.lower()
    for k, v in d.items():
        if str(k).lower() == lk:
            return v
    return None


def _clean_isin(v: Any) -> str | None:
    s = _cell_str(v).upper()
    return s if re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", s) else None


def holdings_from_records(records: list[dict], recipe: Recipe) -> list[Holding]:
    cols = recipe.columns
    if cols is None:
        raise ExtractionError("Recipe has no column map")
    scale = 100.0 if recipe.weight_scale == "fraction" else 1.0
    out = []
    for rec in records:
        name = _cell_str(_ci_get(rec, cols.name))
        weight = to_float(_ci_get(rec, cols.weight))
        if not name or weight is None:
            continue
        out.append(
            Holding(
                name=name,
                weight_pct=weight * scale,
                isin=_clean_isin(_ci_get(rec, cols.isin)),
                country=normalize_country(_cell_str(_ci_get(rec, cols.country)) or None),
                sector=_cell_str(_ci_get(rec, cols.sector)) or None,
            )
        )
    if not out:
        raise ExtractionError("No holdings rows matched the column map")
    return out


def table_to_records(rows: list[list[Any]], recipe: Recipe) -> tuple[list[dict], str, str | None]:
    cols = recipe.columns
    if cols is None:
        raise ExtractionError("Recipe has no column map")
    h = find_header_row(rows, recipe.header_contains, [cols.name, cols.weight])
    header = [_cell_str(v) for v in rows[h]]
    records = [dict(zip(header, row)) for row in rows[h + 1 :] if any(_cell_str(v) for v in row)]
    preamble = "\n".join(" | ".join(_cell_str(v) for v in row if _cell_str(v)) for row in rows[: h + 3])
    as_of = None
    spec = recipe.as_of or ""
    if spec.startswith("cell_right_of:"):
        label = spec.split(":", 1)[1].strip().lower()
        for row in rows[:h + 1]:
            cells = [_cell_str(v).lower() for v in row]
            if label in cells:
                i = cells.index(label)
                for v in row[i + 1 :]:
                    if (d := to_date(v)):
                        as_of = d
                        break
                break
    elif spec.startswith("column:") and records:
        as_of = to_date(_ci_get(records[0], spec.split(":", 1)[1].strip()))
    return records, preamble[:1500], as_of


# ---------------------------------------------------------------- recipe execution


def _parse_payload(content: bytes, fmt: str, recipe: Recipe) -> tuple[list[Holding], str, str | None]:
    if fmt == "json":
        data = json.loads(content)
        records = dot_get(data, recipe.json_holdings_path)
        if not isinstance(records, list):
            raise ExtractionError(f"json_holdings_path {recipe.json_holdings_path!r} is not a list")
        as_of = None
        if recipe.as_of and recipe.as_of.startswith("json:"):
            as_of = to_date(dot_get(data, recipe.as_of.split(":", 1)[1].strip()))
        top = {k: v for k, v in data.items() if not isinstance(v, (list, dict))} if isinstance(data, dict) else {}
        return holdings_from_records(records, recipe), json.dumps(top)[:1500], as_of
    records, preamble, as_of = table_to_records(parse_table(content, fmt), recipe)
    return holdings_from_records(records, recipe), preamble, as_of


def _breakdown(data: Any, path: str | None) -> dict[str, float]:
    items = dot_get(data, path)
    out: dict[str, float] = {}
    if isinstance(items, list):
        for it in items:
            if not isinstance(it, dict):
                continue
            name = normalize_country(it.get("name") or it.get("country") or it.get("label"))
            value = to_float(it.get("value") or it.get("weight") or it.get("percentage"))
            if name and value is not None:
                out[name] = out.get(name, 0.0) + value
    elif isinstance(items, dict):
        for k, v in items.items():
            if (val := to_float(v)) is not None and (name := normalize_country(k)):
                out[name] = val
    return out


async def extract(recipe: Recipe, isin: str, ticker: str, recipe_id: str | None = None) -> FundHoldings:
    """Run a recipe for one fund. Raises ExtractionError with a human-readable reason on failure."""
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    url = fill(recipe.url, isin, ticker)
    page_url = fill(recipe.page_url, isin, ticker)
    countries: dict[str, float] = {}

    if recipe.kind == "static":
        if not recipe.static_holdings:
            raise ExtractionError("static recipe without static_holdings")
        holdings = [Holding(**h.model_dump()) for h in recipe.static_holdings]
        return FundHoldings(
            isin=isin, as_of=date.today().isoformat(), coverage=recipe.coverage, holdings=holdings,
            source_url="static", source_excerpt=recipe.notes, recipe_id=recipe_id, fetched_at=now,
        )

    if recipe.kind == "browser_json":
        if not url:
            raise ExtractionError("browser_json recipe needs url")
        async with browser_session() as b:
            if page_url:
                await b.goto(page_url)
            status, text = await b.fetch_in_page(url)
            if status != 200:
                raise ExtractionError(f"in-page fetch {url} -> HTTP {status}")
            holdings, excerpt, as_of = _parse_payload(text.encode(), "json", recipe)
            if recipe.country_breakdown_url:
                cb_url = fill(recipe.country_breakdown_url, isin, ticker)
                cst, ctext = await b.fetch_in_page(cb_url)
                if cst == 200:
                    countries = _breakdown(json.loads(ctext), recipe.country_breakdown_path)
        return FundHoldings(
            isin=isin, as_of=as_of, coverage=recipe.coverage, holdings=holdings, countries=countries,
            source_url=url, source_excerpt=excerpt, recipe_id=recipe_id, fetched_at=now,
        )

    if recipe.kind == "scrape_link":
        if not page_url or not recipe.link_pattern:
            raise ExtractionError("scrape_link recipe needs page_url and link_pattern")
        page = await http_get(page_url)
        hrefs = re.findall(r"""href=["']([^"']+)["']""", page.text)
        pattern = re.compile(fill(recipe.link_pattern, isin, ticker), re.IGNORECASE)
        match = next((h for h in hrefs if pattern.search(h.replace("&amp;", "&"))), None)
        if not match:
            raise ExtractionError(f"No link on {page_url} matches {recipe.link_pattern!r}")
        url = urljoin(page_url, match.replace("&amp;", "&"))

    if not url:
        raise ExtractionError(f"{recipe.kind} recipe needs url")
    resp = await http_get(normalize_download_url(url))
    fmt = recipe.file_format or sniff_format(resp.content, resp.headers.get("content-type", ""), url)
    if fmt == "html" or sniff_format(resp.content) == "html":
        raise ExtractionError(f"{url} returned an HTML page, not a data file")
    if fmt in ("xls", "xlsx") and sniff_format(resp.content) not in ("xls", "xlsx"):
        fmt = sniff_format(resp.content, resp.headers.get("content-type", ""), url)
    holdings, excerpt, as_of = _parse_payload(resp.content, fmt, recipe)
    if recipe.country_breakdown_url:
        try:
            cb = await http_get(fill(recipe.country_breakdown_url, isin, ticker))
            countries = _breakdown(cb.json(), recipe.country_breakdown_path)
        except Exception:
            countries = {}
    return FundHoldings(
        isin=isin, as_of=as_of, coverage=recipe.coverage, holdings=holdings, countries=countries,
        source_url=url, source_excerpt=excerpt, recipe_id=recipe_id, fetched_at=now,
    )
